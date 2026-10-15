"""
Testes unitários do núcleo do Laudo de Não Conformidades (sem rede e sem banco):
extração de texto, mascaramento LGPD, divisão em partes, interpretação da resposta,
validação de evidência, consolidação, orquestração com provider falso e PDF.
"""

import io
import json

import pytest

from backend.ia import extracao
from backend.ia.extracao import Bloco
from backend.ia import nao_conformidades as nc
from backend.ia.pdf_nao_conformidades import gerar_pdf_nao_conformidades, AVISO_IA
from backend.ia.providers import RespostaIA, ErroProviderIA, criar_provider


class FakeProvider:
    """Provider que devolve respostas pré-definidas, em ordem."""

    def __init__(self, respostas):
        self.respostas = list(respostas)
        self.chamadas = []

    def completar(self, system, user, schema=None):
        self.chamadas.append(user)
        conteudo = self.respostas.pop(0)
        if isinstance(conteudo, Exception):
            raise conteudo
        return RespostaIA(conteudo=conteudo, tokens_entrada=100, tokens_saida=20)


def _item(**kw):
    base = {"ambiente": "Banheiro", "item": "Sifão", "descricao": "Sifão solto com vazamento",
            "categoria": "HIDRAULICA", "gravidade": "MEDIA", "trecho_origem": "sifão solto com vazamento"}
    base.update(kw)
    return base


# --- Extração --------------------------------------------------------------------

def test_blocos_laudo_digital_ignora_ambientes_vazios():
    blocos = extracao.blocos_laudo_digital([
        ("Cozinha", ["Torneira pingando", "  "]),
        ("Sala", []),
    ])
    assert len(blocos) == 1
    assert blocos[0].rotulo == "Ambiente: Cozinha"
    assert blocos[0].texto == "- Torneira pingando"


def test_extrai_texto_de_pdf_e_detecta_digitalizado():
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, "Banheiro social: sifao da pia solto, apresentando vazamento constante.")
    c.showPage()
    c.save()

    paginas = extracao.extrair_paginas_pdf(buf.getvalue())
    assert len(paginas) == 1
    assert "sifao da pia solto" in paginas[0]
    assert not extracao.pdf_parece_digitalizado(paginas)
    assert extracao.pdf_parece_digitalizado(["", "  ", "texto curto"])
    assert extracao.blocos_pdf(paginas)[0].rotulo == "Página 1"


def test_pdf_invalido_gera_erro_de_extracao():
    with pytest.raises(extracao.ErroExtracao):
        extracao.extrair_paginas_pdf(b"isto nao e um pdf")


def test_extrai_texto_de_docx_com_tabela():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Cozinha")
    tabela = doc.add_table(rows=1, cols=2)
    tabela.rows[0].cells[0].text = "Tomada"
    tabela.rows[0].cells[1].text = "Sem funcionamento"
    buf = io.BytesIO()
    doc.save(buf)

    texto = extracao.extrair_texto_docx(buf.getvalue())
    assert "Cozinha" in texto
    assert "Tomada | Sem funcionamento" in texto
    assert extracao.blocos_docx(texto)[0].rotulo == "Seção 1"


def test_mascara_dados_pessoais():
    texto = "Inquilino CPF 123.456.789-09, email joao@exemplo.com, tel (62) 99999-8888. Tomada quebrada."
    mascarado = extracao.mascarar_dados_pessoais(texto)
    assert "123.456.789-09" not in mascarado
    assert "joao@exemplo.com" not in mascarado
    assert "99999-8888" not in mascarado
    assert "[CPF]" in mascarado and "[EMAIL]" in mascarado and "[TELEFONE]" in mascarado
    assert "Tomada quebrada." in mascarado


# --- Divisão em partes -----------------------------------------------------------

def test_divide_em_partes_respeitando_limite():
    blocos = [Bloco(f"Página {i}", "x" * 400) for i in range(10)]
    partes = nc.dividir_em_partes(blocos, max_chars=1000)
    assert len(partes) == 5
    assert all(len(p) <= 1000 for p in partes)
    assert partes[0].startswith("[Página 0]")


def test_bloco_maior_que_limite_e_cortado_por_linhas():
    texto = "\n".join(f"linha {i} " + "y" * 50 for i in range(100))
    partes = nc.dividir_em_partes([Bloco("Página 1", texto)], max_chars=1000)
    assert len(partes) > 1
    assert all(len(p) <= 1100 for p in partes)
    assert "continuação" in partes[1]


# --- Interpretação da resposta ---------------------------------------------------

def test_interpreta_json_com_cercas_e_texto_extra():
    conteudo = "Segue:\n```json\n" + json.dumps({"nao_conformidades": [_item()], "nao_verificados": []}) + "\n```"
    r = nc.interpretar_resposta(conteudo)
    assert len(r.nao_conformidades) == 1


def test_interpreta_json_invalido_gera_erro():
    with pytest.raises(ValueError):
        nc.interpretar_resposta("não há nada aqui")


# --- Evidência e consolidação ----------------------------------------------------

def test_valida_evidencia_exata_aproximada_e_inexistente():
    fonte = nc.normalizar("Banheiro: Sifão da pia SOLTO, com vazamento. Paredes com manchas.")
    assert nc.validar_evidencia("sifão da pia solto, com vazamento", fonte)
    assert nc.validar_evidencia("Sifao da pia solto com vazamentos", fonte)
    assert not nc.validar_evidencia("chuveiro sem funcionamento", fonte)
    assert not nc.validar_evidencia("", fonte)


def test_mescla_duplicatas_mantendo_maior_gravidade():
    a = nc.ItemNC("NAO_CONFORMIDADE", "Banheiro", "Sifão", "Sifão solto com vazamento", "HIDRAULICA", "MEDIA", "t", False)
    b = nc.ItemNC("NAO_CONFORMIDADE", "banheiro", "sifão", "Sifão solto e com vazamento", "HIDRAULICA", "ALTA", "t2", True)
    c = nc.ItemNC("NAO_CONFORMIDADE", "Cozinha", "Tomada", "Tomada sem funcionamento", "ELETRICA", "MEDIA", "t3", True)
    resultado = nc.mesclar_itens([a, b, c])
    assert len(resultado) == 2
    assert resultado[0].gravidade == "ALTA"
    assert resultado[0].evidencia_validada is True


# --- Orquestração ----------------------------------------------------------------

def test_analisar_blocos_valida_evidencia_e_normaliza_campos():
    blocos = [Bloco("Ambiente: Banheiro", "- Sifão solto com vazamento\n- Parede com manchas")]
    resposta = {
        "nao_conformidades": [
            _item(),
            _item(item="Chuveiro", descricao="Chuveiro sem funcionamento", categoria="Elétrica",
                  gravidade="alta", trecho_origem="chuveiro não funciona"),
        ],
        "nao_verificados": [_item(item="Aquecedor", descricao="Não testado", gravidade="CRITICA",
                                  trecho_origem="Parede com manchas")],
    }
    provider = FakeProvider([json.dumps(resposta)])
    resultado = nc.analisar_blocos(blocos, provider)

    assert len(resultado.itens) == 3
    sifao, chuveiro, aquecedor = resultado.itens
    assert sifao.evidencia_validada is True
    assert chuveiro.evidencia_validada is False      # trecho não existe no laudo
    assert chuveiro.categoria == "ELETRICA" and chuveiro.gravidade == "ALTA"
    assert aquecedor.tipo == "NAO_VERIFICADO" and aquecedor.gravidade == "BAIXA"
    assert resultado.tokens_entrada == 100 and resultado.tokens_saida == 20


def test_analisar_blocos_pede_correcao_quando_json_invalido():
    blocos = [Bloco("Página 1", "Tomada sem funcionamento")]
    provider = FakeProvider(["resposta quebrada", json.dumps({"nao_conformidades": [], "nao_verificados": []})])
    resultado = nc.analisar_blocos(blocos, provider)
    assert resultado.itens == []
    assert len(provider.chamadas) == 2
    assert resultado.tokens_entrada == 200


def test_analisar_blocos_falha_se_json_invalido_duas_vezes():
    provider = FakeProvider(["lixo", "mais lixo"])
    with pytest.raises(nc.ErroAnalise):
        nc.analisar_blocos([Bloco("Página 1", "Tomada sem funcionamento")], provider)


def test_analisar_blocos_propaga_erro_do_provider():
    provider = FakeProvider([ErroProviderIA("chave inválida", recuperavel=False)])
    with pytest.raises(nc.ErroAnalise) as exc:
        nc.analisar_blocos([Bloco("Página 1", "Tomada sem funcionamento")], provider)
    assert exc.value.recuperavel is False


def test_analisar_blocos_sem_texto():
    with pytest.raises(nc.ErroAnalise):
        nc.analisar_blocos([Bloco("Página 1", "   ")], FakeProvider([]))


def test_dados_pessoais_nao_sao_enviados_ao_provider():
    provider = FakeProvider([json.dumps({"nao_conformidades": [], "nao_verificados": []})])
    nc.analisar_blocos([Bloco("Página 1", "Proprietário CPF 123.456.789-09. Tomada quebrada.")], provider)
    assert "123.456.789-09" not in provider.chamadas[0]


def test_criar_provider_sem_chave(monkeypatch):
    monkeypatch.delenv("OLLAMA_API_KEY", raising=False)
    with pytest.raises(ErroProviderIA) as exc:
        criar_provider("ollama_cloud", "gemma4:31b")
    assert exc.value.recuperavel is False
    with pytest.raises(ErroProviderIA):
        criar_provider("inexistente", "x")


def test_payload_dos_providers(monkeypatch):
    monkeypatch.setenv("OLLAMA_API_KEY", "k1")
    monkeypatch.setenv("OPENROUTER_API_KEY", "k2")
    ollama = criar_provider("ollama_cloud", "gemma4:31b")
    payload = ollama._payload("sys", "user", nc.SCHEMA_RESPOSTA)
    assert payload["model"] == "gemma4:31b" and payload["format"] == nc.SCHEMA_RESPOSTA
    assert payload["stream"] is False

    openrouter = criar_provider("openrouter", "google/gemma-4-31b-it")
    payload = openrouter._payload("sys", "user", nc.SCHEMA_RESPOSTA)
    assert payload["response_format"]["json_schema"]["schema"] == nc.SCHEMA_RESPOSTA
    assert payload["provider"]["data_collection"] == "deny"
    resp = openrouter._interpretar({"choices": [{"message": {"content": "{}"}}],
                                    "usage": {"prompt_tokens": 5, "completion_tokens": 2}})
    assert resp.tokens_entrada == 5 and resp.tokens_saida == 2


# --- PDF -------------------------------------------------------------------------

def _texto_pdf(buffer) -> str:
    from pypdf import PdfReader
    return "\n".join(p.extract_text() for p in PdfReader(buffer).pages)


def test_gera_pdf_com_itens_e_aviso_de_ia():
    buffer = gerar_pdf_nao_conformidades({
        "codigo_imovel": "CT-123", "endereco": "Rua A, 1 <teste>", "imobiliaria": "Imob X",
        "vistoriador": "Fulano", "data_vistoria": None,
        "itens": [
            {"tipo": "NAO_CONFORMIDADE", "ambiente": "Banheiro", "item": "Sifão",
             "descricao": "Sifão solto com vazamento", "categoria": "HIDRAULICA", "gravidade": "MEDIA"},
            {"tipo": "NAO_VERIFICADO", "ambiente": "Cozinha", "item": "Cooktop",
             "descricao": "Sem gás para teste", "categoria": "GAS", "gravidade": "BAIXA"},
        ],
    })
    texto = _texto_pdf(buffer)
    assert "LAUDO DE NÃO CONFORMIDADES" in texto
    assert AVISO_IA in texto
    assert "Sifão solto com vazamento" in texto
    assert "ITENS NÃO VERIFICADOS" in texto
    assert "Rua A, 1 <teste>" in texto


def test_gera_pdf_sem_itens():
    buffer = gerar_pdf_nao_conformidades({"codigo_imovel": "CT-1", "itens": []})
    assert "Nenhuma não conformidade funcional foi identificada" in _texto_pdf(buffer)
