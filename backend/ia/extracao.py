# backend/ia/extracao.py
"""
Extração de texto das duas fontes de laudo do vistoriador:
  - laudo digital (categorias/ambientes -> itens)
  - arquivo anexado (PDF ou DOCX)

Funções puras (sem banco/rede) para facilitar testes.
"""

import io
import re
import logging
from dataclasses import dataclass
from typing import Iterable, List, Sequence, Tuple

logger = logging.getLogger(__name__)

# Página com menos caracteres que isso é tratada como "sem texto" (imagem/escaneada)
MIN_CHARS_PAGINA_COM_TEXTO = 50


@dataclass
class Bloco:
    """Trecho rotulado do laudo (uma página do PDF ou um ambiente do laudo digital)."""
    rotulo: str
    texto: str


class ErroExtracao(Exception):
    pass


def blocos_laudo_digital(categorias: Iterable[Tuple[str, Sequence[str]]]) -> List[Bloco]:
    """Recebe [(titulo_ambiente, [descricao_item, ...]), ...] na ordem do laudo."""
    blocos = []
    for titulo, descricoes in categorias:
        linhas = [d.strip() for d in descricoes if d and d.strip()]
        if not linhas:
            continue
        texto = "\n".join(f"- {linha}" for linha in linhas)
        blocos.append(Bloco(rotulo=f"Ambiente: {(titulo or '').strip() or 'Sem título'}", texto=texto))
    return blocos


def extrair_paginas_pdf(conteudo: bytes) -> List[str]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(conteudo))
        if reader.is_encrypted:
            reader.decrypt("")
        paginas = []
        for pagina in reader.pages:
            try:
                paginas.append(pagina.extract_text() or "")
            except Exception as e:  # página corrompida não deve derrubar o laudo todo
                logger.warning(f"[IA] Falha ao extrair texto de uma página do PDF: {e}")
                paginas.append("")
        return paginas
    except Exception as e:
        raise ErroExtracao(f"Não foi possível ler o PDF anexado ({e.__class__.__name__}).")


def extrair_texto_docx(conteudo: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(conteudo))
    except Exception as e:
        raise ErroExtracao(f"Não foi possível ler o DOCX anexado ({e.__class__.__name__}).")

    linhas = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for tabela in doc.tables:
        for linha in tabela.rows:
            celulas = []
            for celula in linha.cells:
                texto = celula.text.strip()
                if texto and (not celulas or celulas[-1] != texto):  # células mescladas se repetem
                    celulas.append(texto)
            if celulas:
                linhas.append(" | ".join(celulas))
    return "\n".join(linhas)


def blocos_pdf(paginas: Sequence[str]) -> List[Bloco]:
    return [Bloco(rotulo=f"Página {i}", texto=t.strip())
            for i, t in enumerate(paginas, start=1) if t and t.strip()]


def blocos_docx(texto: str, max_chars: int = 6000) -> List[Bloco]:
    """DOCX não tem páginas: agrupa parágrafos em seções de tamanho razoável."""
    blocos, atual, tamanho = [], [], 0
    for linha in texto.splitlines():
        if tamanho + len(linha) > max_chars and atual:
            blocos.append(Bloco(rotulo=f"Seção {len(blocos) + 1}", texto="\n".join(atual)))
            atual, tamanho = [], 0
        atual.append(linha)
        tamanho += len(linha) + 1
    if atual and "".join(atual).strip():
        blocos.append(Bloco(rotulo=f"Seção {len(blocos) + 1}", texto="\n".join(atual)))
    return blocos


def pdf_parece_digitalizado(paginas: Sequence[str]) -> bool:
    """True quando a maioria das páginas não tem texto extraível (PDF de imagens)."""
    if not paginas:
        return True
    com_texto = sum(1 for p in paginas if len((p or "").strip()) >= MIN_CHARS_PAGINA_COM_TEXTO)
    return com_texto < len(paginas) / 2


_RE_CPF = re.compile(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b")
_RE_CNPJ = re.compile(r"\b\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}\b")
_RE_EMAIL = re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b")
_RE_TELEFONE = re.compile(r"(?:\(?\b\d{2}\)?\s?)?\b9?\d{4}[-\s]?\d{4}\b")


def mascarar_dados_pessoais(texto: str) -> str:
    """Remove identificadores pessoais antes de enviar o texto a um provider externo (LGPD)."""
    texto = _RE_EMAIL.sub("[EMAIL]", texto)
    texto = _RE_CNPJ.sub("[CNPJ]", texto)
    texto = _RE_CPF.sub("[CPF]", texto)
    texto = _RE_TELEFONE.sub("[TELEFONE]", texto)
    return texto
